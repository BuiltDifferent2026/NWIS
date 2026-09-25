import os
import io
import re
import yaml
import openpyxl
from typing import Dict, Any, List, Optional, Tuple

from app.models.ddr_models import (
    DDRNormalizedData,
    SourceMetadata,
    WellFields,
    FormationFields,
    DrillingParameterFields,
    BitBHAFields,
    TimeAccountingFields,
    NarrativeFields,
    ExtractedField,
    DetectedSection,
    ExtractionStatus,
    ConfidenceLevel
)
from app.services.ddr_normalizer import (
    clean_text,
    normalize_label,
    extract_number_and_unit,
    extract_npt_from_narrative,
    parse_colon_value,
    compute_relative_depth
)
from app.services.ddr_validation import validate_ddr_extraction

# Load mapping config
CONFIG_PATH = os.path.join(os.path.dirname(__file__), "..", "config", "ddr_template_mapping.yaml")

def load_template_mapping() -> Dict[str, Any]:
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)

class DDRSectionParser:
    def __init__(self, mapping_config: Optional[Dict[str, Any]] = None):
        self.config = mapping_config or load_template_mapping()
        self.sections_config = self.config.get("sections", [])
        self.fields_config = self.config.get("fields", {})

    def parse_workbook(self, file_content: bytes, filename: str = "DDR.xlsx") -> DDRNormalizedData:
        """Parse an uploaded DDR Excel workbook bytes."""
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        sheet = wb.active
        return self.parse_worksheet(sheet, filename=filename)

    def parse_file_path(self, file_path: str) -> DDRNormalizedData:
        """Parse an Excel file from disk."""
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        filename = os.path.basename(file_path)
        return self.parse_worksheet(sheet, filename=filename)

    def _build_merged_cell_map(self, sheet) -> Dict[Tuple[int, int], Tuple[int, int]]:
        """Map every (row, col) inside a merged range to the range's top-left (min_row, min_col)."""
        merged_map = {}
        for rng in sheet.merged_cells.ranges:
            top_left = (rng.min_row, rng.min_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged_map[(r, c)] = top_left
        return merged_map

    def _get_cell_value_and_coord(self, sheet, r: int, c: int, merged_map: Dict) -> Tuple[Any, str]:
        """Resolve value from top-left cell if in a merged range, and return coordinate string."""
        if (r, c) in merged_map:
            actual_r, actual_c = merged_map[(r, c)]
        else:
            actual_r, actual_c = r, c
        val = sheet.cell(row=actual_r, column=actual_c).value
        col_letter = openpyxl.utils.get_column_letter(actual_c)
        return val, f"{col_letter}{actual_r}"

    def parse_worksheet(self, sheet, filename: str = "DDR.xlsx") -> DDRNormalizedData:
        merged_map = self._build_merged_cell_map(sheet)
        max_r = min(sheet.max_row, 100)
        max_c = min(sheet.max_column, 26)

        # Index all sheet cells: (r, c) -> (clean_val, coord, raw_val)
        cell_matrix: Dict[Tuple[int, int], Dict[str, Any]] = {}
        row_text_index: Dict[int, List[Tuple[int, str, str, Any]]] = {}

        for r in range(1, max_r + 1):
            row_text_index[r] = []
            for c in range(1, max_c + 1):
                raw_val, coord = self._get_cell_value_and_coord(sheet, r, c, merged_map)
                val_text = clean_text(raw_val)
                cell_matrix[(r, c)] = {
                    "val": val_text,
                    "coord": coord,
                    "raw": raw_val,
                    "r": r,
                    "c": c
                }
                if val_text:
                    row_text_index[r].append((c, val_text, coord, raw_val))

        # Detect Sections
        detected_sections = self._detect_sections(sheet, max_r, max_c, cell_matrix)

        # Extract Fields
        extracted_dict: Dict[str, ExtractedField] = {}
        for field_id, field_cfg in self.fields_config.items():
            extracted_dict[field_id] = self._extract_single_field(
                field_id, field_cfg, cell_matrix, max_r, max_c
            )

        # Secondary Pass: NPT regex check from summary line if either npt_24h_h or monthly_cumulative_npt_h is missing
        needs_npt_daily = extracted_dict.get("npt_24h_h") is None or extracted_dict["npt_24h_h"].status != ExtractionStatus.FOUND
        needs_npt_cum = extracted_dict.get("monthly_cumulative_npt_h") is None or extracted_dict["monthly_cumulative_npt_h"].status != ExtractionStatus.FOUND
        if needs_npt_daily or needs_npt_cum:
            self._fallback_npt_extraction(cell_matrix, extracted_dict, max_r, max_c)

        # Build Normalized Data Structure
        norm_data = self._assemble_normalized_data(
            extracted_dict, detected_sections, sheet.title, filename
        )

        # Generate a lightweight 2D sheet preview grid for UI Screen 2 (rows 1-45, cols A-F)
        preview_grid = []
        for r in range(1, min(max_r + 1, 46)):
            row_obj = {"row": r, "cells": []}
            for c in range(1, min(max_c + 1, 7)):
                cell_info = cell_matrix.get((r, c), {})
                col_letter = openpyxl.utils.get_column_letter(c)
                row_obj["cells"].append({
                    "col": col_letter,
                    "coord": f"{col_letter}{r}",
                    "val": cell_info.get("val", "")
                })
            preview_grid.append(row_obj)
        norm_data.sheet_preview_grid = preview_grid

        # Run Validation Summary
        norm_data.validation_summary = validate_ddr_extraction(norm_data)

        return norm_data

    def _detect_sections(
        self, sheet, max_r: int, max_c: int, cell_matrix: Dict
    ) -> List[DetectedSection]:
        """Identify section boundary rows based on keywords and headers."""
        found_sections: List[DetectedSection] = []
        
        # Candidate section rows
        for sec in self.sections_config:
            sec_id = sec["id"]
            sec_name = sec["name"]
            color = sec.get("color", "blue")
            keywords = [k.lower() for k in sec.get("keywords", [])]

            best_r = -1
            matched_lbl = ""
            for r in range(1, max_r + 1):
                for c in range(1, min(max_c + 1, 8)):
                    val = normalize_label(cell_matrix.get((r, c), {}).get("val", ""))
                    if not val:
                        continue
                    for kw in keywords:
                        if kw == val or (len(kw) > 6 and kw in val):
                            best_r = r
                            matched_lbl = cell_matrix.get((r, c), {}).get("val", "")
                            break
                    if best_r != -1:
                        break
                if best_r != -1:
                    break

            if best_r != -1:
                found_sections.append(DetectedSection(
                    id=sec_id,
                    name=sec_name,
                    color=color,
                    start_row=best_r,
                    end_row=best_r + 3,  # default span
                    matched_label=matched_lbl,
                    sample_preview=f"Section identified at row {best_r}"
                ))

        # Sort by start_row and adjust end_rows
        found_sections.sort(key=lambda s: s.start_row)
        for i in range(len(found_sections)):
            if i < len(found_sections) - 1:
                found_sections[i].end_row = found_sections[i + 1].start_row - 1
            else:
                found_sections[i].end_row = max_r

        return found_sections

    def _extract_single_field(
        self,
        field_id: str,
        field_cfg: Dict[str, Any],
        cell_matrix: Dict,
        max_r: int,
        max_c: int
    ) -> ExtractedField:
        aliases = [normalize_label(a) for a in field_cfg.get("aliases", [])]
        aliases.sort(key=len, reverse=True)
        expected_unit = field_cfg.get("expected_unit")
        direction = field_cfg.get("direction", "right")
        confidence_str = field_cfg.get("confidence", "STRUCTURED_HIGH")
        confidence = (
            ConfidenceLevel.STRUCTURED_MEDIUM 
            if confidence_str == "STRUCTURED_MEDIUM" 
            else ConfidenceLevel.STRUCTURED_HIGH
        )

        # 1. Exact Match Pass
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                cell_item = cell_matrix.get((r, c), {})
                raw_cell_text = cell_item.get("val", "")
                if not raw_cell_text:
                    continue

                norm_cell_text = normalize_label(raw_cell_text)

                for alias in aliases:
                    if norm_cell_text == alias:
                        return self._resolve_field_value(
                            r, c, raw_cell_text, field_cfg, direction, expected_unit, confidence, cell_matrix, max_c
                        )

                # Check colon pattern exact match (e.g. 'Well: WELL_A' or 'Hrs Lost: 2.5')
                if ":" in raw_cell_text:
                    lbl_part, inline_val = parse_colon_value(raw_cell_text)
                    norm_lbl_part = normalize_label(lbl_part)
                    for alias in aliases:
                        if norm_lbl_part == alias:
                            return self._resolve_field_value(
                                r, c, raw_cell_text, field_cfg, direction, expected_unit, confidence, cell_matrix, max_c
                            )

        # 2. Prefix / Substring Match Pass
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                cell_item = cell_matrix.get((r, c), {})
                raw_cell_text = cell_item.get("val", "")
                if not raw_cell_text:
                    continue

                norm_cell_text = normalize_label(raw_cell_text)

                for alias in aliases:
                    if len(alias) >= 5 and norm_cell_text.startswith(alias):
                        return self._resolve_field_value(
                            r, c, raw_cell_text, field_cfg, direction, expected_unit, confidence, cell_matrix, max_c
                        )

        return ExtractedField(
            value=None,
            unit=expected_unit,
            source_cell=None,
            status=ExtractionStatus.MISSING,
            confidence=confidence
        )

    def _resolve_field_value(
        self,
        r: int,
        c: int,
        label_text: str,
        field_cfg: Dict[str, Any],
        direction: str,
        expected_unit: Optional[str],
        confidence: ConfidenceLevel,
        cell_matrix: Dict,
        max_c: int
    ) -> ExtractedField:
        # Check inline colon value first
        if ":" in label_text:
            _, inline_val = parse_colon_value(label_text)
            if inline_val:
                return self._build_extracted_field(
                    inline_val,
                    cell_matrix.get((r, c), {}).get("coord", f"R{r}C{c}"),
                    expected_unit,
                    confidence
                )

        # Direction: BELOW (for narrative fields)
        if direction == "below":
            below_cell = cell_matrix.get((r + 1, c), {})
            below_val = below_cell.get("val", "")
            below_coord = below_cell.get("coord", f"R{r+1}C{c}")
            if below_val:
                return ExtractedField(
                    value=below_val,
                    unit=expected_unit,
                    source_cell=below_coord,
                    status=ExtractionStatus.FOUND,
                    confidence=ConfidenceLevel.STRUCTURED_MEDIUM,
                    raw_value=below_val
                )
            else:
                return ExtractedField(
                    value=None,
                    unit=expected_unit,
                    source_cell=below_coord,
                    status=ExtractionStatus.MISSING,
                    confidence=ConfidenceLevel.STRUCTURED_MEDIUM
                )

        # Direction: RIGHT (or adjacent)
        # Search next cell(s) in the same row
        for target_c in range(c + 1, min(max_c + 1, c + 4)):
            candidate_cell = cell_matrix.get((r, target_c), {})
            val = candidate_cell.get("val")
            coord = candidate_cell.get("coord", f"R{r}C{target_c}")
            
            # If candidate cell is the same merged cell as label, skip
            if coord == cell_matrix.get((r, c), {}).get("coord"):
                continue

            if val is not None and str(val).strip():
                return self._build_extracted_field(val, coord, expected_unit, confidence)

        # If no value found to the right, check row below if appropriate
        if direction in ("right_or_inline", "right"):
            below_cell = cell_matrix.get((r + 1, c), {})
            val = below_cell.get("val")
            if val is not None and str(val).strip():
                coord = below_cell.get("coord", f"R{r+1}C{c}")
                return self._build_extracted_field(val, coord, expected_unit, confidence)

        return ExtractedField(
            value=None,
            unit=expected_unit,
            source_cell=cell_matrix.get((r, c), {}).get("coord"),
            status=ExtractionStatus.MISSING,
            confidence=confidence
        )

    def _build_extracted_field(
        self,
        raw_val: Any,
        source_cell: str,
        expected_unit: Optional[str],
        confidence: ConfidenceLevel
    ) -> ExtractedField:
        str_val = clean_text(raw_val)
        if not str_val:
            return ExtractedField(
                value=None,
                unit=expected_unit,
                source_cell=source_cell,
                status=ExtractionStatus.MISSING,
                confidence=confidence,
                raw_value=raw_val
            )

        # Parse numeric if expected unit is numeric
        if expected_unit in ("m", "in", "t", "ft-lb", "rpm", "gpm", "psi", "ppg", "sec/qt", "cc/30min", "hrs"):
            num, unit = extract_number_and_unit(str_val, default_unit=expected_unit)
            if num is not None:
                return ExtractedField(
                    value=num,
                    unit=unit or expected_unit,
                    source_cell=source_cell,
                    status=ExtractionStatus.FOUND,
                    confidence=confidence,
                    raw_value=raw_val
                )
            else:
                return ExtractedField(
                    value=str_val,
                    unit=expected_unit,
                    source_cell=source_cell,
                    status=ExtractionStatus.NEEDS_REVIEW,
                    confidence=confidence,
                    raw_value=raw_val
                )

        return ExtractedField(
            value=str_val,
            unit=expected_unit if expected_unit != "text" else None,
            source_cell=source_cell,
            status=ExtractionStatus.FOUND,
            confidence=confidence,
            raw_value=raw_val
        )

    def _fallback_npt_extraction(self, cell_matrix: Dict, extracted_dict: Dict, max_r: int, max_c: int):
        """Check for summary line like 'NPT in 24 hrs: X hrs, Cumm monthly NPT: Y hrs'."""
        for (r, c), item in cell_matrix.items():
            text = item.get("val", "")
            if "npt in 24 hrs" in text.lower() or "cumm monthly npt" in text.lower():
                npt_24h, cum_npt = extract_npt_from_narrative(text)
                coord = item.get("coord", f"R{r}C{c}")
                if npt_24h is not None and (extracted_dict.get("npt_24h_h") is None or extracted_dict["npt_24h_h"].status != ExtractionStatus.FOUND):
                    extracted_dict["npt_24h_h"] = ExtractedField(
                        value=npt_24h,
                        unit="hrs",
                        source_cell=coord,
                        status=ExtractionStatus.FOUND,
                        confidence=ConfidenceLevel.STRUCTURED_HIGH,
                        raw_value=text
                    )
                if cum_npt is not None and (extracted_dict.get("monthly_cumulative_npt_h") is None or extracted_dict["monthly_cumulative_npt_h"].status != ExtractionStatus.FOUND):
                    extracted_dict["monthly_cumulative_npt_h"] = ExtractedField(
                        value=cum_npt,
                        unit="hrs",
                        source_cell=coord,
                        status=ExtractionStatus.FOUND,
                        confidence=ConfidenceLevel.STRUCTURED_HIGH,
                        raw_value=text
                    )
                break

    def _assemble_normalized_data(
        self,
        extracted: Dict[str, ExtractedField],
        detected_sections: List[DetectedSection],
        sheet_name: str,
        filename: str
    ) -> DDRNormalizedData:
        # Calculate relative depth if both present_depth and formation_top are available
        present_depth_val = extracted.get("present_depth_md_m", ExtractedField()).value
        top_md_val = extracted.get("top_md_m", ExtractedField()).value
        
        rel_depth_val = compute_relative_depth(
            float(present_depth_val) if isinstance(present_depth_val, (int, float)) else None,
            float(top_md_val) if isinstance(top_md_val, (int, float)) else None
        )

        relative_depth_field = ExtractedField(
            value=rel_depth_val,
            unit="m",
            source_cell=f"{extracted.get('present_depth_md_m', ExtractedField()).source_cell or ''} - {extracted.get('top_md_m', ExtractedField()).source_cell or ''}".strip(" -"),
            status=ExtractionStatus.FOUND if rel_depth_val is not None else ExtractionStatus.MISSING,
            confidence=ConfidenceLevel.STRUCTURED_HIGH
        )

        source_meta = SourceMetadata(
            document_type="DDR",
            sheet_name=sheet_name,
            filename=filename,
            report_date=extracted.get("report_date", ExtractedField())
        )

        well_fields = WellFields(
            well_id=extracted.get("well_id", ExtractedField()),
            basin=extracted.get("basin", ExtractedField()),
            location=extracted.get("location", ExtractedField()),
            rig_name=extracted.get("rig_name", ExtractedField()),
            present_depth_md_m=extracted.get("present_depth_md_m", ExtractedField()),
            progress_24h_m=extracted.get("progress_24h_m", ExtractedField()),
            present_operation=extracted.get("present_operation", ExtractedField()),
            hole_size_in=extracted.get("hole_size_in", ExtractedField())
        )

        formation_fields = FormationFields(
            top_md_m=extracted.get("top_md_m", ExtractedField()),
            current_formation=extracted.get("current_formation", ExtractedField()),
            relative_depth_m=relative_depth_field
        )

        drilling_params = DrillingParameterFields(
            wob_t=extracted.get("wob_t", ExtractedField()),
            torque_on_bottom_ft_lb=extracted.get("torque_on_bottom_ft_lb", ExtractedField()),
            torque_off_bottom_ft_lb=extracted.get("torque_off_bottom_ft_lb", ExtractedField()),
            rpm=extracted.get("rpm", ExtractedField()),
            flow_rate_gpm=extracted.get("flow_rate_gpm", ExtractedField()),
            spp_psi=extracted.get("spp_psi", ExtractedField()),
            mud_weight_ppg=extracted.get("mud_weight_ppg", ExtractedField()),
            viscosity_sec_qt=extracted.get("viscosity_sec_qt", ExtractedField()),
            fluid_loss_cc_30min=extracted.get("fluid_loss_cc_30min", ExtractedField())
        )

        bit_bha = BitBHAFields(
            bit_serial_no=extracted.get("bit_serial_no", ExtractedField()),
            bit_type_nozzle=extracted.get("bit_type_nozzle", ExtractedField()),
            make_iadc=extracted.get("make_iadc", ExtractedField()),
            bit_hours_24h=extracted.get("bit_hours_24h", ExtractedField()),
            bit_hours_cum=extracted.get("bit_hours_cum", ExtractedField()),
            bha_length_m=extracted.get("bha_length_m", ExtractedField()),
            bha_weight_t=extracted.get("bha_weight_t", ExtractedField())
        )

        time_accounting = TimeAccountingFields(
            hours_lost=extracted.get("hours_lost", ExtractedField()),
            shutdown_reason=extracted.get("shutdown_reason", ExtractedField()),
            operating_time_h=extracted.get("operating_time_h", ExtractedField()),
            force_majeure_h=extracted.get("force_majeure_h", ExtractedField()),
            repairing_time_h=extracted.get("repairing_time_h", ExtractedField()),
            standby_day_time_h=extracted.get("standby_day_time_h", ExtractedField()),
            non_operating_time_h=extracted.get("non_operating_time_h", ExtractedField()),
            ilm_time_h=extracted.get("ilm_time_h", ExtractedField()),
            npt_24h_h=extracted.get("npt_24h_h", ExtractedField()),
            monthly_cumulative_npt_h=extracted.get("monthly_cumulative_npt_h", ExtractedField())
        )

        narrative_fields = NarrativeFields(
            operation_details=extracted.get("operation_details", ExtractedField()),
            next_operation=extracted.get("next_operation", ExtractedField()),
            remarks=extracted.get("remarks", ExtractedField())
        )

        return DDRNormalizedData(
            source=source_meta,
            well=well_fields,
            formation=formation_fields,
            drilling_parameters=drilling_params,
            bit_bha=bit_bha,
            time_accounting=time_accounting,
            narrative=narrative_fields,
            detected_sections=detected_sections
        )
