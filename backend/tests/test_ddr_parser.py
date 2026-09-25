import io
import openpyxl
import pytest

from app.services.ddr_normalizer import (
    extract_number_and_unit,
    extract_npt_from_narrative,
    compute_relative_depth,
    parse_colon_value,
    normalize_label
)
from app.services.ddr_parser import DDRSectionParser
from app.models.ddr_models import ExtractionStatus, ConfidenceLevel, DataQuality

def create_mock_ddr_workbook() -> bytes:
    """Generate a clean synthetic DDR workbook in memory for unit testing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DDR_Daily"

    # Merge cell A1:B2 for header test
    ws.merge_cells("A1:B2")
    ws["A1"] = "WELL CONTEXT SECTION"

    # Well Header with merged cells and colons
    ws["C1"] = "Basin"
    ws.merge_cells("D1:F1")
    ws["D1"] = "ASSAM_ARAKAN"

    ws["C2"] = "Rig"
    ws["D2"] = "RIG_X"

    ws["C3"] = "Well: WELL_A"

    ws["A4"] = "Operation Date"
    ws["B4"] = "2026-09-24"

    ws["A5"] = "Present Depth (M)"
    ws["B5"] = "2,145 m"

    ws["C5"] = "Formation Top(m)"
    ws["D5"] = "2,080 m"

    ws["A6"] = "Formation Drilled"
    ws["B6"] = "FORMATION_ALPHA"

    # Drilling parameters
    ws["A8"] = "Drilling Parameters"
    ws["A9"] = "Mud Wt. (ppg)"
    ws["B9"] = "10.2 ppg"

    ws["C9"] = "WOB (T)"
    ws["D9"] = "12 T"

    ws["E9"] = "Torque (On Bottom) (ft-lb)"
    ws["F9"] = "11,000 ft-lb"

    ws["A10"] = "RPM"
    ws["B10"] = "120"

    ws["C10"] = "GPM"
    ws["D10"] = "550"

    ws["E10"] = "SPP (psi)"
    ws["F10"] = "2,800 psi"

    # Narrative
    ws["A12"] = "Operation Details"
    ws["A13"] = "Drilling 8.5 inch hole smoothly from 2069m to 2145m with controlled ROP."

    # Time accounting & NPT
    ws["A15"] = "Time Lost due to Shut Down"
    ws["A16"] = "Hrs Lost: 2.5"
    ws["B16"] = "Reason: Rig pump liner repair"

    ws["A18"] = "NPT in 24 hrs: 2.5 hrs, Cumm monthly NPT: 14.0 hrs"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ─── 1. Merged-Cell Resolution Test ───
def test_merged_cell_resolution():
    wb_bytes = create_mock_ddr_workbook()
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes), data_only=True)
    ws = wb.active

    parser = DDRSectionParser()
    merged_map = parser._build_merged_cell_map(ws)

    # A1:B2 was merged -> B2 should resolve to (1, 1)
    assert (2, 2) in merged_map
    assert merged_map[(2, 2)] == (1, 1)
    val, coord = parser._get_cell_value_and_coord(ws, 2, 2, merged_map)
    assert val == "WELL CONTEXT SECTION"
    assert coord == "A1"

    # D1:F1 was merged -> F1 should resolve to D1
    assert (1, 6) in merged_map
    assert merged_map[(1, 6)] == (1, 4)
    val, coord = parser._get_cell_value_and_coord(ws, 1, 6, merged_map)
    assert val == "ASSAM_ARAKAN"
    assert coord == "D1"

# ─── 2. Label Alias Matching Test ───
def test_label_alias_matching():
    parser = DDRSectionParser()
    aliases = parser.fields_config["present_depth_md_m"]["aliases"]
    
    # Matches various synonyms
    assert normalize_label("Present Depth (M)") in [normalize_label(a) for a in aliases]
    assert normalize_label("current depth (m)") in [normalize_label(a) for a in aliases]
    assert normalize_label("Hole Depth") in [normalize_label(a) for a in aliases]

    # Test colon parsing
    lbl, val = parse_colon_value("Well: WELL_A")
    assert lbl == "Well"
    assert val == "WELL_A"

    lbl2, val2 = parse_colon_value("Hrs Lost: 2.5")
    assert lbl2 == "Hrs Lost"
    assert val2 == "2.5"

# ─── 3. Numeric Parsing With Units Test ───
def test_numeric_parsing_with_units():
    # Number with comma and unit
    num, unit = extract_number_and_unit("2,145 m", default_unit="m")
    assert num == 2145.0
    assert unit == "m"

    # Float with unit
    num2, unit2 = extract_number_and_unit("10.2 ppg", default_unit="ppg")
    assert num2 == 10.2
    assert unit2 == "ppg"

    # Number with ft-lb unit
    num3, unit3 = extract_number_and_unit("11,000 ft-lb", default_unit="ft-lb")
    assert num3 == 11000.0
    assert unit3 == "ft-lb"

    # Integer without unit
    num4, unit4 = extract_number_and_unit("120", default_unit="rpm")
    assert num4 == 120.0
    assert unit4 == "rpm"

    # Non-numeric text returns None
    num5, _ = extract_number_and_unit("FORMATION_ALPHA")
    assert num5 is None

# ─── 4. NPT Summary Narrative Parsing Test ───
def test_npt_summary_parsing():
    text1 = "NPT in 24 hrs: 2.5 hrs, Cumm monthly NPT: 14.0 hrs"
    daily, cumm = extract_npt_from_narrative(text1)
    assert daily == 2.5
    assert cumm == 14.0

    text2 = "NPT in 24 hrs: 0 hrs, Cumm monthly NPT:0 hrs"
    daily2, cumm2 = extract_npt_from_narrative(text2)
    assert daily2 == 0.0
    assert cumm2 == 0.0

# ─── 5. Formation-Relative Depth Calculation Test ───
def test_formation_relative_depth_calculation():
    # Both depths available
    rel = compute_relative_depth(2145.0, 2080.0)
    assert rel == 65.0

    # Decimals
    rel2 = compute_relative_depth(2165.4, 2100.0)
    assert rel2 == 65.4

    # None if either is missing
    assert compute_relative_depth(None, 2080.0) is None
    assert compute_relative_depth(2145.0, None) is None

# ─── 6. End-to-End Workbook Extraction & Validation Test ───
def test_full_workbook_extraction_and_validation():
    wb_bytes = create_mock_ddr_workbook()
    parser = DDRSectionParser()
    extracted = parser.parse_workbook(wb_bytes, filename="test_mock_ddr.xlsx")

    # Check extracted values
    assert extracted.well.well_id.value == "WELL_A"
    assert extracted.well.well_id.status == ExtractionStatus.FOUND

    assert extracted.well.present_depth_md_m.value == 2145.0
    assert extracted.well.present_depth_md_m.unit == "m"
    assert extracted.well.present_depth_md_m.source_cell == "B5"

    assert extracted.formation.top_md_m.value == 2080.0
    assert extracted.formation.current_formation.value == "FORMATION_ALPHA"
    assert extracted.formation.relative_depth_m.value == 65.0

    assert extracted.drilling_parameters.mud_weight_ppg.value == 10.2
    assert extracted.drilling_parameters.wob_t.value == 12.0
    assert extracted.drilling_parameters.torque_on_bottom_ft_lb.value == 11000.0
    assert extracted.drilling_parameters.spp_psi.value == 2800.0

    assert extracted.time_accounting.hours_lost.value == 2.5
    assert extracted.time_accounting.npt_24h_h.value == 2.5
    assert extracted.time_accounting.monthly_cumulative_npt_h.value == 14.0

    assert extracted.narrative.operation_details.value.startswith("Drilling 8.5 inch hole")
    assert extracted.narrative.operation_details.confidence == ConfidenceLevel.STRUCTURED_MEDIUM

    # Check validation summary
    assert extracted.validation_summary.fields_found > 10
    assert extracted.validation_summary.data_quality in (DataQuality.HIGH, DataQuality.MEDIUM)
    assert any("Required drilling parameters" in m for m in extracted.validation_summary.validation_messages)
